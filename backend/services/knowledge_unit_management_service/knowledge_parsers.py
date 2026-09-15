"""知识单元管理服务 · 文档解析与文本切片。

对应《技术方案设计文档》5.3 的第 2、3 项职责：

    2. 格式解析：按 file_type 分派解析器（PDF / Markdown / Word / TXT）
    3. 文本切片：按段落与长度阈值切分

设计约束：

* 支持格式严格限定为 2.9.3 声明的四类（PDF、Markdown、Word、TXT），
  其余一律不接受，由 :func:`detect_file_type` 返回 ``None`` 表达"不支持"；
* PDF 与 Word 的解析必须依赖第三方库（pypdf、python-docx），因此这两个解析器
  **按需导入**——只有真正解析该格式时才要求环境里装了对应库，
  Markdown 与 TXT 场景不会被无关依赖拖死；
* 切片只做"按段落 + 长度阈值"两件事。文档未规定切片重叠窗口，因此不设重叠。
"""

from __future__ import annotations

import io
import re
from pathlib import Path

__all__ = [
    "SUPPORTED_FILE_TYPES",
    "CHUNK_MAX_CHARS",
    "CHUNK_MIN_CHARS",
    "detect_file_type",
    "parse_document",
    "split_into_chunks",
]

# 扩展名 -> knowledge_units.file_type 取值。取值集合与 models.enums.FileType 一致，
# 也与 2.9.3 知识导入中心声明的四种支持格式一一对应。
# Word 只接受 .docx：.doc 是二进制老格式，pypdf/python-docx 这类库都不支持，
# 需求写的是 "Word" 而非具体版本，此处按可解析的 .docx 落位。
SUPPORTED_FILE_TYPES = {
    ".pdf": "pdf",
    ".md": "md",
    ".markdown": "md",
    ".docx": "docx",
    ".txt": "txt",
}

# 切片长度阈值。需求只写了"按段落与长度阈值切分"但未给出具体数值，
# 因此这里给出默认值并暴露为函数参数，便于调整而不必改代码逻辑。
CHUNK_MAX_CHARS = 800  # 单个切片的最大字符数
CHUNK_MIN_CHARS = 100  # 尾块小于该值时视为碎片

# 段落分隔：一个或多个空行
_PARAGRAPH_SPLIT_RE = re.compile(r"\n\s*\n")
# 句子分隔：中英文常见句末标点（用于超长段落的二次切分）
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[。！？；.!?;])\s*")
# 连续空行压缩
_BLANK_LINES_RE = re.compile(r"\n{3,}")


def detect_file_type(file_name: str) -> str | None:
    """根据文件名后缀判断知识单元的文件类型。

    :param file_name: 上传文件名，允许带目录前缀（2.9.3 支持整个目录批量导入）。
    :return: ``pdf`` / ``md`` / ``docx`` / ``txt``；不在支持范围内时返回 ``None``，
        由调用方翻译成接口响应里的 ``rejected`` 条目（见 8.4 的 ``unsupported_format``）。
    """
    # 目录批量导入时前端可能带上相对路径，这里只取最后一段文件名
    suffix = Path(file_name).suffix.lower()
    return SUPPORTED_FILE_TYPES.get(suffix)


def parse_document(file_name: str, data: bytes) -> str:
    """把上传的原始字节解析成纯文本正文。

    :param file_name: 文件名，用于决定走哪个解析器。
    :param data: 文件原始字节。
    :return: 解析后的纯文本；解析器可能返回空串（例如扫描版 PDF 抽不出文字），
        这种情况由调用方决定如何处理，本函数不擅自编造内容。
    :raise ValueError: 文件类型不在 2.9.3 声明的四类之内。
    """
    file_type = detect_file_type(file_name)
    if file_type is None:
        raise ValueError(f"不支持的文件类型：{file_name}")
    if file_type == "pdf":
        return _parse_pdf(data)
    if file_type == "md":
        return _parse_markdown(data)
    if file_type == "docx":
        return _parse_docx(data)
    return _parse_txt(data)


def _parse_pdf(data: bytes) -> str:
    """解析 PDF：逐页抽取文本后按页拼接。

    扫描件 PDF 抽不出文字层，这里如实返回空串，不做 OCR（需求未提及 OCR）。
    """
    # 按需导入：不解析 PDF 的场景无需安装 pypdf
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        # extract_text() 对纯图片页会返回 None，统一转成空串保证拼接不报错
        pages.append(page.extract_text() or "")
    return _normalize("\n\n".join(pages))


def _parse_markdown(data: bytes) -> str:
    """解析 Markdown：Markdown 本身即纯文本，解码后保留标记符号原文。

    不剥离 ``#``、``*`` 等标记 —— 需求只要求"解析文档内容"，
    保留原标记可避免丢失标题层级信息，也不影响后续切片。
    """
    return _normalize(_decode_text(data))


def _parse_docx(data: bytes) -> str:
    """解析 Word（.docx）：抽取全部段落与表格单元格文本。

    只取文字，忽略图片与样式（需求未要求保留样式）。
    """
    # 按需导入：不解析 Word 的场景无需安装 python-docx
    from docx import Document

    document = Document(io.BytesIO(data))
    # 先按正文段落顺序取文字
    parts = [p.text for p in document.paragraphs]
    # 再补充表格内的文字：表格内容同样是知识正文的一部分，漏掉会丢信息
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return _normalize("\n\n".join(part for part in parts if part and part.strip()))


def _parse_txt(data: bytes) -> str:
    """解析 TXT：解码后直接作为正文。"""
    return _normalize(_decode_text(data))


def _decode_text(data: bytes) -> str:
    """把字节按常见编码解码成字符串。

    先试 UTF-8（当前主流），再试 GB18030（兼容 Windows 上常见的中文 txt），
    两种都失败时按 UTF-8 容错解码，避免个别坏字符导致整篇导入失败。
    """
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    """统一换行符并压缩连续空行，保证后续切片的分段判断稳定。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return _BLANK_LINES_RE.sub("\n\n", text).strip()


def _split_long_paragraph(paragraph: str, max_chars: int) -> list[str]:
    """把超过长度阈值的单个段落切成多块。

    策略：先按句末标点切句，再按阈值把句子累加成块；若单个句子本身就超长
    （例如无标点的长表格行），则退化为按字符数硬切，保证输出块不超过阈值。
    """
    pieces: list[str] = []
    buffer = ""
    for sentence in _SENTENCE_SPLIT_RE.split(paragraph):
        if not sentence:
            continue
        # 单句超长：先把已累积的内容收尾，再把这一句按阈值硬切
        while len(sentence) > max_chars:
            if buffer:
                pieces.append(buffer)
                buffer = ""
            pieces.append(sentence[:max_chars])
            sentence = sentence[max_chars:]
        # 累加当前句子，超出阈值则先把已累积的内容收尾
        if len(buffer) + len(sentence) > max_chars:
            pieces.append(buffer)
            buffer = sentence
        else:
            buffer += sentence
    if buffer:
        pieces.append(buffer)
    return pieces


def split_into_chunks(
    text: str,
    max_chars: int = CHUNK_MAX_CHARS,
    min_chars: int = CHUNK_MIN_CHARS,
) -> list[str]:
    """把正文按"段落 + 长度阈值"切成检索粒度。

    4.3 与 5.3 的口径是：**每个独立导入的文档作为一个知识单元**，
    而切片只是该单元下的检索粒度，因此本函数只负责产出切片列表，
    不负责决定知识单元的边界（那是服务主模块的事）。

    三步走：

    1. 按空行拆段落，超长段落按句切分，保证没有任何一块超过 ``max_chars``；
    2. 相邻块顺序累加，累加到再加一块就会超阈值时收尾，减少碎片数量；
    3. 尾块若小于 ``min_chars`` 且合并后仍不超阈值，则并入前一块，
       避免产生检索价值极低的碎片；若不满足合并条件则原样保留。

    :param text: 已解析出的纯文本正文。
    :param max_chars: 单块最大字符数。
    :param min_chars: 尾块碎片判定阈值。
    :return: 切片列表；正文为空时返回空列表。
    """
    normalized = _normalize(text)
    if not normalized:
        return []

    # 第一步：拆段落，并保证单块不超阈值
    blocks: list[str] = []
    for paragraph in _PARAGRAPH_SPLIT_RE.split(normalized):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) <= max_chars:
            blocks.append(paragraph)
        else:
            blocks.extend(_split_long_paragraph(paragraph, max_chars))

    # 第二步：相邻块累加，尽量填满但不越过阈值
    chunks: list[str] = []
    buffer = ""
    for block in blocks:
        if not buffer:
            buffer = block
        elif len(buffer) + len(block) + 1 <= max_chars:
            buffer = f"{buffer}\n{block}"
        else:
            chunks.append(buffer)
            buffer = block
    if buffer:
        chunks.append(buffer)

    # 第三步：尾块过短则并入前一块，前提是合并后仍不超阈值
    if len(chunks) >= 2 and len(chunks[-1]) < min_chars:
        if len(chunks[-2]) + len(chunks[-1]) + 1 <= max_chars:
            chunks[-2] = f"{chunks[-2]}\n{chunks[-1]}"
            chunks.pop()

    return chunks
